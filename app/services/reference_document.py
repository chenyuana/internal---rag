from __future__ import annotations

from hashlib import sha256
from io import BytesIO
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from app.core.exceptions import AppError
from app.schemas.chat import ReferenceDocument, ReferenceDocumentParsed
from app.schemas.retrieval import ChunkMetadata, Citation, NormalizedQuery, SelectedChunk
from app.services.evidence_text import overlap_score

MAX_REFERENCE_BYTES = 10 * 1024 * 1024
MAX_REFERENCE_CHARS = 20_000
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}
MAX_AUXILIARY_CHUNKS = 8
MAX_AUXILIARY_CHUNK_CHARS = 1_800


def parse_reference_document(filename: str, content: bytes) -> ReferenceDocumentParsed:
    """Extract bounded text without storing or publishing the uploaded file."""
    safe_name = Path(filename or "reference-document").name
    extension = Path(safe_name).suffix.casefold()
    if extension not in SUPPORTED_EXTENSIONS:
        raise AppError(
            code="REFERENCE_DOCUMENT_UNSUPPORTED",
            message="参考文档仅支持 PDF、DOCX、TXT 和 Markdown。",
            status_code=415,
        )
    if not content:
        raise AppError(
            code="REFERENCE_DOCUMENT_EMPTY",
            message="参考文档为空。",
            status_code=400,
        )
    if len(content) > MAX_REFERENCE_BYTES:
        raise AppError(
            code="REFERENCE_DOCUMENT_TOO_LARGE",
            message="参考文档不能超过 10 MB。",
            status_code=413,
        )

    try:
        if extension == ".pdf":
            reader = PdfReader(BytesIO(content))
            text = "\n\n".join((page.extract_text() or "").strip() for page in reader.pages)
        elif extension == ".docx":
            document = Document(BytesIO(content))
            blocks = [paragraph.text.strip() for paragraph in document.paragraphs]
            for table in document.tables:
                blocks.extend(
                    " | ".join(cell.text.strip() for cell in row.cells)
                    for row in table.rows
                )
            text = "\n".join(block for block in blocks if block)
        else:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("gb18030")
    except Exception as exc:
        raise AppError(
            code="REFERENCE_DOCUMENT_PARSE_FAILED",
            message="无法读取参考文档，请确认文件未损坏且未加密。",
            status_code=422,
        ) from exc

    normalized = "\n".join(line.rstrip() for line in text.splitlines()).strip()
    if not normalized:
        raise AppError(
            code="REFERENCE_DOCUMENT_NO_TEXT",
            message="参考文档中没有可提取的文字；扫描版 PDF 暂不支持作为格式参考。",
            status_code=422,
        )
    truncated = len(normalized) > MAX_REFERENCE_CHARS
    return ReferenceDocumentParsed(
        name=safe_name,
        text=normalized[:MAX_REFERENCE_CHARS],
        truncated=truncated,
    )


def build_auxiliary_evidence(
    document: ReferenceDocument,
    query: NormalizedQuery,
    *,
    citation_start: int,
) -> tuple[list[SelectedChunk], list[Citation]]:
    """Turn a request-local document into relevant, citable evidence chunks."""
    raw_chunks = _split_text(document.text)
    if query.query_type in {"summary", "procedure"}:
        chosen = raw_chunks[:MAX_AUXILIARY_CHUNKS]
    else:
        ranked = sorted(
            enumerate(raw_chunks),
            key=lambda item: (
                -overlap_score(query.normalized_query, f"{document.name}\n{item[1]}"),
                item[0],
            ),
        )
        chosen = [text for _, text in ranked[:MAX_AUXILIARY_CHUNKS]]

    digest = sha256(document.text.encode("utf-8")).hexdigest()[:16]
    document_id = f"auxiliary-{digest}"
    document_name = f"{document.name}（临时辅助资料）"
    selected: list[SelectedChunk] = []
    citations: list[Citation] = []
    for offset, text in enumerate(chosen):
        citation_id = f"C{citation_start + offset}"
        chunk_id = f"{document_id}-chunk-{offset + 1}"
        evidence_text = f"文档：{document.name}\n{text}"
        metadata = ChunkMetadata(
            document_name=document_name,
            source_path="temporary-upload",
            chunk_index=offset,
            document_type="temporary_auxiliary",
            content_type="text",
            status="effective",
        )
        selected.append(
            SelectedChunk(
                citation_id=citation_id,
                chunk_id=chunk_id,
                document_id=document_id,
                dataset_id="temporary-auxiliary",
                text=evidence_text,
                metadata=metadata,
                hybrid_score=1.0,
                vector_score=1.0,
                keyword_score=1.0,
                rerank_score=1.0,
            )
        )
        citations.append(
            Citation(
                citation_id=citation_id,
                chunk_id=chunk_id,
                document_id=document_id,
                document_name=document_name,
                source_path="temporary-upload",
                quote=text,
            )
        )
    return selected, citations


def _split_text(text: str) -> list[str]:
    paragraphs = [part.strip() for part in text.replace("\r", "").split("\n\n") if part.strip()]
    chunks: list[str] = []
    buffer = ""
    for paragraph in paragraphs:
        if len(paragraph) > MAX_AUXILIARY_CHUNK_CHARS:
            if buffer:
                chunks.append(buffer)
                buffer = ""
            chunks.extend(
                paragraph[index : index + MAX_AUXILIARY_CHUNK_CHARS]
                for index in range(0, len(paragraph), MAX_AUXILIARY_CHUNK_CHARS)
            )
            continue
        candidate = f"{buffer}\n\n{paragraph}".strip()
        if buffer and len(candidate) > MAX_AUXILIARY_CHUNK_CHARS:
            chunks.append(buffer)
            buffer = paragraph
        else:
            buffer = candidate
    if buffer:
        chunks.append(buffer)
    return chunks or [text[:MAX_AUXILIARY_CHUNK_CHARS]]
