from __future__ import annotations

import hashlib
import html as html_lib
import re
import statistics
import xml.etree.ElementTree as ET
import zipfile
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docx
from docx import Document
from docx.document import Document as DocxDocument
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph

from app.core.exceptions import AppError
from app.ingestion.pipeline import (
    AssetRecord,
    BlockRecord,
    PageRecord,
    ParsedDocument,
    build_chunks,
    compact_chars,
    sha256_file,
    stable_id,
    table_rows_from_html,
    table_to_semantic_text,
)

DOCX_MIME_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
DOCX_MAX_UNCOMPRESSED_BYTES = 1024 * 1024 * 1024
DOCX_MAX_SINGLE_PART_BYTES = 256 * 1024 * 1024
_TABLE_CAPTION_RE = re.compile(r"^(?:表|Table)\s*[A-Za-z0-9一二三四五六七八九十]", re.I)
_HEADING_STYLE_RE = re.compile(r"^(?:Heading|标题)\s*([1-9])$", re.I)
_VML_IMAGE_DATA_TAG = "{urn:schemas-microsoft-com:vml}imagedata"


@dataclass(slots=True)
class _RenderedCell:
    text: str
    column: int
    colspan: int = 1
    rowspan: int = 1


def validate_docx_package(path: Path) -> tuple[int, int]:
    """Validate the OOXML package and return (part_count, uncompressed_bytes)."""

    try:
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            required = {"[Content_Types].xml", "word/document.xml"}
            if not required.issubset(names):
                raise ValueError("required Word OOXML parts are missing")
            infos = package.infolist()
            total = sum(item.file_size for item in infos)
            largest = max((item.file_size for item in infos), default=0)
            if total > DOCX_MAX_UNCOMPRESSED_BYTES:
                raise ValueError("uncompressed OOXML package exceeds the safety limit")
            if largest > DOCX_MAX_SINGLE_PART_BYTES:
                raise ValueError("an OOXML package part exceeds the safety limit")
            return len(infos), total
    except (OSError, zipfile.BadZipFile, ValueError) as exc:
        raise AppError(
            code="DOCUMENT_SIGNATURE_MISMATCH",
            message="The uploaded .docx file is not a valid Word OOXML document.",
            status_code=415,
            details={"extension": ".docx", "reason": str(exc)},
        ) from exc


def _normalized_text(value: str) -> str:
    lines = [" ".join(line.replace("\x00", "").split()) for line in value.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _visible_text_segments(element: Any) -> list[str]:
    """Read visible OOXML text, including hyperlinks/content controls/insertions."""

    segments = [""]
    deleted = {qn("w:del"), qn("w:moveFrom")}
    text_tags = {qn("w:t"), qn("m:t"), qn("w:delText")}

    def append(value: str) -> None:
        segments[-1] += value

    def walk(node: Any) -> None:
        if node.tag in deleted:
            return
        if node.tag in text_tags:
            if node.tag != qn("w:delText") and node.text:
                append(node.text)
            return
        if node.tag == qn("w:tab"):
            append("\t")
            return
        if node.tag in {qn("w:cr"), qn("w:br")}:
            break_type = node.get(qn("w:type"))
            if node.tag == qn("w:br") and break_type == "page":
                segments.append("")
            else:
                append("\n")
            return
        if node.tag == qn("w:lastRenderedPageBreak"):
            segments.append("")
            return
        if node.tag in {qn("w:instrText"), qn("w:fldChar")}:
            return
        for child in node:
            walk(child)

    walk(element)
    return [_normalized_text(value) for value in segments]


def _paragraph_heading_level(paragraph: Paragraph) -> int | None:
    style_name = str(getattr(paragraph.style, "name", "") or "").strip()
    style_id = str(getattr(paragraph.style, "style_id", "") or "").strip()
    if style_name.casefold() == "title" or style_id.casefold() == "title":
        return 0
    for candidate in (style_name, style_id.replace("Heading", "Heading ")):
        match = _HEADING_STYLE_RE.match(candidate.strip())
        if match:
            return int(match.group(1))
    paragraph_properties = paragraph._p.pPr
    if paragraph_properties is not None:
        outline = paragraph_properties.find(qn("w:outlineLvl"))
        if outline is not None:
            try:
                return int(outline.get(qn("w:val"))) + 1
            except (TypeError, ValueError):
                pass
    return None


def _paragraph_break_before(paragraph: Paragraph) -> bool:
    properties = paragraph._p.pPr
    if properties is None:
        return False
    marker = properties.find(qn("w:pageBreakBefore"))
    if marker is None:
        return False
    return str(marker.get(qn("w:val"), "1")).casefold() not in {"0", "false", "off"}


def _effective_num_pr(paragraph: Paragraph) -> Any | None:
    properties = paragraph._p.pPr
    if properties is not None and properties.numPr is not None:
        return properties.numPr
    style = getattr(paragraph, "style", None)
    style_properties = getattr(getattr(style, "element", None), "pPr", None)
    return getattr(style_properties, "numPr", None)


class _ListNumbering:
    def __init__(self, document: DocxDocument) -> None:
        self._numbering = document.part.numbering_part.element
        self._counters: dict[tuple[int, int], int] = {}
        self._style_counters: defaultdict[str, int] = defaultdict(int)

    def _format(self, num_id: int, level: int) -> tuple[str, int]:
        num = next(
            (
                item
                for item in self._numbering.findall(qn("w:num"))
                if int(item.get(qn("w:numId"), -1)) == num_id
            ),
            None,
        )
        if num is None:
            return "decimal", 1
        abstract_id_node = num.find(qn("w:abstractNumId"))
        if abstract_id_node is None:
            return "decimal", 1
        abstract_id = int(abstract_id_node.get(qn("w:val"), -1))
        abstract = next(
            (
                item
                for item in self._numbering.findall(qn("w:abstractNum"))
                if int(item.get(qn("w:abstractNumId"), -1)) == abstract_id
            ),
            None,
        )
        if abstract is None:
            return "decimal", 1
        level_node = next(
            (
                item
                for item in abstract.findall(qn("w:lvl"))
                if int(item.get(qn("w:ilvl"), -1)) == level
            ),
            None,
        )
        if level_node is None:
            return "decimal", 1
        format_node = level_node.find(qn("w:numFmt"))
        start_node = level_node.find(qn("w:start"))
        number_format = (
            str(format_node.get(qn("w:val"), "decimal"))
            if format_node is not None
            else "decimal"
        )
        try:
            start = int(start_node.get(qn("w:val"), 1)) if start_node is not None else 1
        except (TypeError, ValueError):
            start = 1
        return number_format, start

    def prefix(self, paragraph: Paragraph) -> str:
        num_pr = _effective_num_pr(paragraph)
        style_name = str(getattr(paragraph.style, "name", "") or "")
        if num_pr is None:
            if "bullet" in style_name.casefold() or "项目符号" in style_name:
                return "- "
            if "number" in style_name.casefold() or "编号" in style_name:
                self._style_counters[style_name] += 1
                return f"{self._style_counters[style_name]}. "
            return ""
        try:
            num_id = int(num_pr.numId.val)
            level = int(num_pr.ilvl.val) if num_pr.ilvl is not None else 0
        except (AttributeError, TypeError, ValueError):
            return "- "
        number_format, start = self._format(num_id, level)
        if number_format == "bullet":
            return "  " * level + "- "
        key = (num_id, level)
        self._counters[key] = self._counters.get(key, start - 1) + 1
        for counter_key in list(self._counters):
            if counter_key[0] == num_id and counter_key[1] > level:
                self._counters.pop(counter_key, None)
        return "  " * level + f"{self._counters[key]}. "


def _table_header_rows(table: Table) -> int:
    count = 0
    for row in table.rows:
        properties = row._tr.trPr
        if properties is None or properties.find(qn("w:tblHeader")) is None:
            break
        count += 1
    return max(1, count)


def _cell_span(cell_element: Any) -> int:
    properties = cell_element.tcPr
    marker = properties.find(qn("w:gridSpan")) if properties is not None else None
    try:
        return max(1, int(marker.get(qn("w:val"), 1))) if marker is not None else 1
    except (TypeError, ValueError):
        return 1


def _cell_vertical_merge(cell_element: Any) -> str | None:
    properties = cell_element.tcPr
    marker = properties.find(qn("w:vMerge")) if properties is not None else None
    if marker is None:
        return None
    value = marker.get(qn("w:val"))
    return "restart" if value == "restart" else "continue"


def _cell_text(cell_element: Any, table: Table) -> str:
    try:
        cell = _Cell(cell_element, table)
        parts = [
            text
            for paragraph in cell.paragraphs
            for text in _visible_text_segments(paragraph._p)
            if text
        ]
        return "\n".join(parts).strip()
    except (AttributeError, TypeError, ValueError):
        return "\n".join(
            text for text in _visible_text_segments(cell_element) if text
        ).strip()


def _table_to_html(table: Table, *, header_rows: int) -> tuple[str, int, int]:
    rendered_rows: list[list[_RenderedCell]] = []
    active_vertical: dict[int, _RenderedCell] = {}
    merged_cell_count = 0
    nested_table_count = 0
    for row_element in table._tbl.tr_lst:
        row_cells: list[_RenderedCell] = []
        next_vertical: dict[int, _RenderedCell] = {}
        extended: set[int] = set()
        column = 0
        for cell_element in row_element.tc_lst:
            colspan = _cell_span(cell_element)
            vertical_merge = _cell_vertical_merge(cell_element)
            if vertical_merge == "continue":
                origin = active_vertical.get(column)
                if origin is not None:
                    identity = id(origin)
                    if identity not in extended:
                        origin.rowspan += 1
                        merged_cell_count += 1
                        extended.add(identity)
                    for offset in range(colspan):
                        next_vertical[column + offset] = origin
                    column += colspan
                    continue
            text = _cell_text(cell_element, table)
            cell = _RenderedCell(text=text, column=column, colspan=colspan)
            row_cells.append(cell)
            if colspan > 1:
                merged_cell_count += 1
            if vertical_merge == "restart":
                for offset in range(colspan):
                    next_vertical[column + offset] = cell
            column += colspan
            nested_table_count += len(cell_element.findall(qn("w:tbl")))
        rendered_rows.append(row_cells)
        active_vertical = next_vertical

    html_rows: list[str] = []
    for row_index, cells in enumerate(rendered_rows):
        tag = "th" if row_index < header_rows else "td"
        parts: list[str] = []
        current_column = 0
        for cell in cells:
            if cell.column > current_column:
                parts.extend(f"<{tag}></{tag}>" for _ in range(cell.column - current_column))
            attributes = ""
            if cell.rowspan > 1:
                attributes += f' rowspan="{cell.rowspan}"'
            if cell.colspan > 1:
                attributes += f' colspan="{cell.colspan}"'
            parts.append(
                f"<{tag}{attributes}>{html_lib.escape(cell.text)}</{tag}>"
            )
            current_column = cell.column + cell.colspan
        html_rows.append("<tr>" + "".join(parts) + "</tr>")
    return "<table>" + "".join(html_rows) + "</table>", merged_cell_count, nested_table_count


def _image_relationship_ids(element: Any) -> list[str]:
    ids: list[str] = []
    for node in element.iter():
        if node.tag == qn("a:blip"):
            relationship_id = node.get(qn("r:embed"))
        elif node.tag == _VML_IMAGE_DATA_TAG:
            relationship_id = node.get(qn("r:id"))
        else:
            continue
        if relationship_id and relationship_id not in ids:
            ids.append(relationship_id)
    return ids


def _image_alt_texts(element: Any) -> list[str]:
    values: list[str] = []
    for node in element.iter():
        if node.tag not in {qn("wp:docPr"), qn("pic:cNvPr")}:
            continue
        for attribute in ("descr", "title", "name"):
            value = str(node.get(attribute) or "").strip()
            if value and value not in values and not re.fullmatch(r"Picture \d+", value, re.I):
                values.append(value)
    return values


def _supplemental_notes(path: Path, member_name: str) -> list[str]:
    try:
        with zipfile.ZipFile(path) as package:
            if member_name not in package.namelist():
                return []
            root = ET.fromstring(package.read(member_name))
    except (OSError, zipfile.BadZipFile, ET.ParseError):
        return []
    item_tag = qn("w:footnote") if "footnotes" in member_name else qn("w:endnote")
    notes: list[str] = []
    for item in root.findall(item_tag):
        try:
            note_id = int(item.get(qn("w:id"), -1))
        except (TypeError, ValueError):
            note_id = -1
        if note_id < 0:
            continue
        text = "\n".join(value for value in _visible_text_segments(item) if value)
        if text:
            notes.append(f"[{note_id}] {text}")
    return notes


class DocxParser:
    name = "native-docx"
    version = docx.__version__
    supported_extensions = frozenset({".docx"})

    def parse(self, path: Path) -> ParsedDocument:
        package_part_count, uncompressed_bytes = validate_docx_package(path)
        try:
            document = Document(path)
        except (OSError, ValueError, KeyError, PackageNotFoundError) as exc:
            raise AppError(
                code="DOCUMENT_PARSE_FAILED",
                message="The Word document could not be parsed.",
                status_code=422,
                details={"error_type": type(exc).__name__},
            ) from exc

        source_hash = sha256_file(path)
        document_id = stable_id("document", source_hash)
        pages: dict[int, PageRecord] = {}
        blocks: list[BlockRecord] = []
        assets: list[AssetRecord] = []
        page_text: defaultdict[int, list[str]] = defaultdict(list)
        page_number = 1
        heading_stack: list[tuple[int, str]] = []
        pending_table_caption: str | None = None
        table_count = 0
        image_count = 0
        image_without_alt_count = 0
        unsupported_drawing_count = 0
        nested_table_count = 0
        list_item_count = 0
        numbering = _ListNumbering(document)

        def ensure_page(number: int) -> PageRecord:
            if number not in pages:
                pages[number] = PageRecord(
                    page_number=number,
                    raw_text="",
                    cleaned_text="",
                    route="docx_native",
                    page_type="document",
                    layout_features={"pagination_mode": "word_breaks"},
                )
            return pages[number]

        def current_section() -> list[str]:
            return [text for _, text in heading_stack]

        def add_block(
            *,
            block_type: str,
            text: str,
            number: int,
            **metadata: Any,
        ) -> BlockRecord | None:
            normalized = text.strip()
            if not normalized:
                return None
            block = BlockRecord(
                block_id=stable_id(document_id, number, len(blocks), normalized),
                block_type=block_type,
                text=normalized,
                page_number=number,
                section_path=current_section(),
                source_page_start=number,
                source_page_end=number,
                **metadata,
            )
            blocks.append(block)
            page_text[number].append(normalized)
            rich = {
                "block_type": block.block_type,
                "text": block.text,
                "bbox": None,
                "confidence": 1.0,
                "table_html": block.table_html,
                "asset_id": block.asset_id,
                "asset_ids": list(block.asset_ids),
                "source_page_start": number,
                "source_page_end": number,
                "table_id": block.table_id,
                "table_title": block.table_title,
                "table_rows": block.table_rows,
                "table_header_rows": block.table_header_rows,
                "table_row_pages": block.table_row_pages,
            }
            ensure_page(number).rich_blocks.append(rich)
            return block

        def extract_assets(element: Any, number: int) -> tuple[list[str], list[str]]:
            nonlocal image_count, image_without_alt_count
            relationship_ids = _image_relationship_ids(element)
            if not relationship_ids:
                return [], []
            alt_texts = _image_alt_texts(element)
            asset_ids: list[str] = []
            descriptions: list[str] = []
            for relationship_id in relationship_ids:
                part = document.part.related_parts.get(relationship_id)
                blob = getattr(part, "blob", None)
                if not isinstance(blob, bytes):
                    continue
                image_count += 1
                filename = Path(str(getattr(part, "partname", "image.bin"))).name
                mime_type = str(getattr(part, "content_type", "application/octet-stream"))
                description = next(iter(alt_texts), "")
                if not description:
                    image_without_alt_count += 1
                    description = f"Word 文档图片：{filename}"
                asset_id = f"p{number:04d}-docx-image-{image_count:02d}"
                assets.append(
                    AssetRecord(
                        asset_id=asset_id,
                        page_number=number,
                        asset_type="figure",
                        bbox=[0.0, 0.0, 0.0, 0.0],
                        filename=filename,
                        mime_type=mime_type,
                        content=blob,
                        caption=description,
                        description=description,
                    )
                )
                asset_ids.append(asset_id)
                descriptions.append(description)
            return asset_ids, descriptions

        ensure_page(page_number)
        for element in document.element.body.iterchildren():
            if element.tag == qn("w:p"):
                paragraph = Paragraph(element, document)
                if _paragraph_break_before(paragraph) and (blocks or page_text[page_number]):
                    page_number += 1
                    ensure_page(page_number)
                segments = _visible_text_segments(element)
                asset_ids, image_descriptions = extract_assets(element, page_number)
                drawing_count = sum(
                    1 for node in element.iter() if node.tag == qn("w:drawing")
                )
                unsupported_drawing_count += max(0, drawing_count - len(asset_ids))
                heading_level = _paragraph_heading_level(paragraph)
                list_prefix = numbering.prefix(paragraph)
                if list_prefix:
                    list_item_count += 1
                for segment_index, segment in enumerate(segments):
                    if segment_index:
                        page_number += 1
                        ensure_page(page_number)
                    text = (list_prefix + segment).strip() if segment else ""
                    if not text:
                        continue
                    style_name = str(getattr(paragraph.style, "name", "") or "")
                    is_table_caption = bool(
                        _TABLE_CAPTION_RE.match(text)
                        and (
                            style_name.casefold() == "caption"
                            or style_name == "题注"
                            or text.startswith(("表", "Table"))
                        )
                    )
                    if is_table_caption:
                        if pending_table_caption:
                            add_block(
                                block_type="paragraph",
                                text=pending_table_caption,
                                number=page_number,
                            )
                        pending_table_caption = text
                        continue
                    if pending_table_caption:
                        add_block(
                            block_type="paragraph",
                            text=pending_table_caption,
                            number=page_number,
                        )
                        pending_table_caption = None
                    if heading_level is not None:
                        heading_stack = [
                            (level, value)
                            for level, value in heading_stack
                            if level < heading_level
                        ]
                        heading_stack.append((heading_level, text))
                        add_block(
                            block_type="document_heading",
                            text=text,
                            number=page_number,
                        )
                    else:
                        add_block(
                            block_type="paragraph",
                            text=text,
                            number=page_number,
                            asset_ids=asset_ids,
                            asset_id=next(iter(asset_ids), None),
                        )
                if asset_ids and not any(segments):
                    for asset_id, description in zip(
                        asset_ids,
                        image_descriptions,
                        strict=False,
                    ):
                        add_block(
                            block_type="figure",
                            text=description,
                            number=page_number,
                            asset_id=asset_id,
                            asset_ids=[asset_id],
                        )
                continue

            if element.tag != qn("w:tbl"):
                continue
            table = Table(element, document)
            table_count += 1
            header_rows = _table_header_rows(table)
            table_html, merged_cell_count, nested_count = _table_to_html(
                table,
                header_rows=header_rows,
            )
            nested_table_count += nested_count
            rows = table_rows_from_html(table_html)
            asset_ids, _ = extract_assets(element, page_number)
            if not rows:
                continue
            title = pending_table_caption or ""
            pending_table_caption = None
            table_id = stable_id(
                document_id,
                "docx-table",
                table_count,
                page_number,
                title,
            )
            table_text = table_to_semantic_text(
                rows,
                title=title,
                header_rows=header_rows,
            )
            add_block(
                block_type="table",
                text=table_text,
                number=page_number,
                table_html=table_html,
                table_id=table_id,
                table_title=title or None,
                table_rows=rows,
                table_header_rows=header_rows,
                table_row_pages=[page_number] * len(rows),
                asset_id=next(iter(asset_ids), None),
                asset_ids=asset_ids,
            )
            ensure_page(page_number).page_type = "document_table"
            ensure_page(page_number).table_hints.append(title or f"Word 表格 {table_count}")
            if merged_cell_count:
                ensure_page(page_number).layout_features.setdefault(
                    "merged_table_cell_count",
                    0,
                )
                ensure_page(page_number).layout_features["merged_table_cell_count"] += (
                    merged_cell_count
                )

        if pending_table_caption:
            add_block(
                block_type="paragraph",
                text=pending_table_caption,
                number=page_number,
            )

        footnotes = _supplemental_notes(path, "word/footnotes.xml")
        endnotes = _supplemental_notes(path, "word/endnotes.xml")
        for label, notes in (("脚注", footnotes), ("尾注", endnotes)):
            if not notes:
                continue
            heading_stack = [(1, label)]
            add_block(
                block_type="document_heading",
                text=label,
                number=page_number,
            )
            for note in notes:
                add_block(block_type="paragraph", text=note, number=page_number)

        ordered_pages = [pages[number] for number in sorted(pages)]
        for page in ordered_pages:
            text = "\n\n".join(page_text[page.page_number]).strip()
            page.raw_text = text
            page.cleaned_text = text
            page.raw_char_count = len(compact_chars(text))
            page.cleaned_char_count = page.raw_char_count
            page.image_count = sum(
                asset.page_number == page.page_number for asset in assets
            )

        chunks = build_chunks(document_id, blocks)
        chunk_lengths = [len(chunk.text) for chunk in chunks]
        text_pages = sum(page.raw_char_count > 0 for page in ordered_pages)
        coverage = text_pages / len(ordered_pages) if ordered_pages else 0.0
        quality_gates: list[dict[str, str]] = []
        if not blocks or not chunks:
            quality_gates.append(
                {
                    "status": "fail",
                    "gate": "empty_document",
                    "message": "Word 文档中没有可索引的正文、表格或图片说明。",
                }
            )
        if unsupported_drawing_count:
            quality_gates.append(
                {
                    "status": "warn",
                    "gate": "word_drawing_semantics",
                    "message": (
                        "检测到无法直接提取为图片的 Word 绘图/图表对象，需人工抽检："
                        f"{unsupported_drawing_count}"
                    ),
                }
            )
        if image_without_alt_count:
            quality_gates.append(
                {
                    "status": "warn",
                    "gate": "word_image_semantics",
                    "message": (
                        "Word 图片缺少替代文字，图片已保留但语义需要人工补充："
                        f"{image_without_alt_count}"
                    ),
                }
            )
        if nested_table_count:
            quality_gates.append(
                {
                    "status": "warn",
                    "gate": "word_nested_table",
                    "message": f"检测到嵌套 Word 表格，需抽检结构：{nested_table_count}",
                }
            )
        if not quality_gates:
            quality_gates.append(
                {
                    "status": "pass",
                    "gate": "native_docx",
                    "message": "Word 正文、标题层级和表格结构已完成原生解析。",
                }
            )

        normalized_text = compact_chars("\n".join(page.raw_text for page in ordered_pages))
        normalized_text_hash = (
            hashlib.sha256(normalized_text.encode("utf-8")).hexdigest()
            if normalized_text
            else None
        )
        warnings = sum(gate["status"] == "warn" for gate in quality_gates)
        failures = sum(gate["status"] == "fail" for gate in quality_gates)
        route = "docx_review" if warnings or failures else "docx_native"
        qa = {
            "page_count": len(ordered_pages),
            "pagination_mode": "word_explicit_and_last_rendered_breaks",
            "text_page_count": text_pages,
            "content_page_count": len(ordered_pages),
            "text_coverage": round(coverage, 4),
            "raw_char_count": sum(page.raw_char_count for page in ordered_pages),
            "cleaned_char_count": sum(page.cleaned_char_count for page in ordered_pages),
            "heading_count": sum(block.block_type == "document_heading" for block in blocks),
            "table_count": table_count,
            "structured_table_count": sum(block.block_type == "table" for block in blocks),
            "table_pages": sorted(
                {block.page_number for block in blocks if block.block_type == "table"}
            ),
            "image_count": image_count,
            "image_without_alt_count": image_without_alt_count,
            "unsupported_drawing_count": unsupported_drawing_count,
            "nested_table_count": nested_table_count,
            "list_item_count": list_item_count,
            "footnote_count": len(footnotes),
            "endnote_count": len(endnotes),
            "package_part_count": package_part_count,
            "package_uncompressed_bytes": uncompressed_bytes,
            "chunk_count": len(chunks),
            "chunk_char_min": min(chunk_lengths, default=0),
            "chunk_char_median": int(statistics.median(chunk_lengths)) if chunk_lengths else 0,
            "chunk_char_max": max(chunk_lengths, default=0),
            "quality_gates": quality_gates,
            "parser_name": self.name,
            "parser_trace": [
                {
                    "parser": self.name,
                    "status": "completed",
                    "paragraph_count": len(document.paragraphs),
                    "table_count": table_count,
                    "logical_page_count": len(ordered_pages),
                }
            ],
        }
        return ParsedDocument(
            source_path=path,
            source_hash=source_hash,
            normalized_text_hash=normalized_text_hash,
            document_id=document_id,
            pages=ordered_pages,
            route=route,
            repeated_margin_lines=[],
            removed_margin_line_count=0,
            blocks=blocks,
            chunks=chunks,
            qa=qa,
            assets=assets,
        )
