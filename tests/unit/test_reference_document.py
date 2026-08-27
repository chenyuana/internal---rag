from __future__ import annotations

from io import BytesIO

from docx import Document

from app.schemas.chat import ReferenceDocument
from app.schemas.retrieval import NormalizedQuery
from app.services.reference_document import (
    MAX_REFERENCE_CHARS,
    build_auxiliary_evidence,
    parse_reference_document,
)


def test_parse_text_reference_document() -> None:
    result = parse_reference_document("示例.md", "# 标题\n\n第一节内容".encode())

    assert result.name == "示例.md"
    assert result.text.startswith("# 标题")
    assert result.truncated is False


def test_parse_docx_reference_document_includes_paragraphs_and_tables() -> None:
    document = Document()
    document.add_heading("报告标题", level=1)
    document.add_paragraph("摘要内容")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "字段"
    table.cell(0, 1).text = "说明"
    stream = BytesIO()
    document.save(stream)

    result = parse_reference_document("报告.docx", stream.getvalue())

    assert "报告标题" in result.text
    assert "字段 | 说明" in result.text


def test_parse_reference_document_truncates_model_payload() -> None:
    result = parse_reference_document("long.txt", ("段落内容\n" * 5000).encode())

    assert len(result.text) == MAX_REFERENCE_CHARS
    assert result.truncated is True


def test_reference_document_parse_endpoint(client) -> None:
    response = client.post(
        "/api/v1/chat/reference-documents/parse",
        files={"file": ("style.txt", "一、概述\n二、结论".encode(), "text/plain")},
    )

    assert response.status_code == 200
    assert response.json()["name"] == "style.txt"


def test_build_auxiliary_evidence_is_citable_and_request_local() -> None:
    chunks, citations = build_auxiliary_evidence(
        ReferenceDocument(
            name="补充说明.md",
            text="## 飞行要求\n\n最大飞行高度为 120 米。\n\n## 维护要求\n\n每月检查电池。",
        ),
        NormalizedQuery(
            original_query="最大飞行高度是多少？",
            normalized_query="最大飞行高度是多少？",
            query_type="fact",
        ),
        citation_start=4,
    )

    assert chunks[0].citation_id == "C4"
    assert chunks[0].dataset_id == "temporary-auxiliary"
    assert chunks[0].metadata.document_name == "补充说明.md（临时辅助资料）"
    assert "120 米" in chunks[0].text
    assert citations[0].citation_id == "C4"
    assert citations[0].source_path == "temporary-upload"
