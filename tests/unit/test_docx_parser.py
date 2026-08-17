from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from PIL import Image

from app.core.exceptions import AppError
from app.ingestion.parsers import DocxParser, ParserRegistry


def _word_document(path: Path) -> None:
    document = Document()
    document.add_heading("无人机测试报告", level=0)
    document.add_heading("1 适用范围", level=1)
    document.add_paragraph("本文件规定了试验对象、环境条件和验收准则。")
    document.add_paragraph("记录环境温度", style="List Bullet")
    document.add_paragraph("检查供电状态", style="List Bullet")
    caption = document.add_paragraph("表1 试验参数", style="Caption")
    caption.alignment = 1
    table = document.add_table(rows=3, cols=3)
    header_marker = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(header_marker)
    for cell, value in zip(table.rows[0].cells, ["参数", "下限", "上限"], strict=True):
        cell.text = value
    table.cell(1, 0).text = "温度"
    table.cell(1, 1).text = "-20℃"
    table.cell(1, 2).text = "55℃"
    table.cell(2, 0).text = "共同要求"
    merged = table.cell(2, 1).merge(table.cell(2, 2))
    merged.text = "设备应正常工作"
    document.add_page_break()
    document.add_heading("1.1 验收方法", level=2)
    document.add_paragraph("逐项核对测试记录并确认结果。")
    document.save(path)


def test_docx_parser_preserves_headings_lists_tables_and_page_breaks(
    tmp_path: Path,
) -> None:
    source = tmp_path / "无人机测试报告.docx"
    _word_document(source)

    parsed = DocxParser().parse(source)

    assert parsed.route == "docx_native"
    assert parsed.qa["parser_name"] == "native-docx"
    assert parsed.qa["page_count"] == 2
    assert parsed.qa["heading_count"] == 3
    assert parsed.qa["list_item_count"] == 2
    assert parsed.qa["structured_table_count"] == 1
    assert parsed.qa["quality_gates"] == [
        {
            "status": "pass",
            "gate": "native_docx",
            "message": "Word 正文、标题层级和表格结构已完成原生解析。",
        }
    ]
    assert "- 记录环境温度" in parsed.pages[0].raw_text
    assert "逐项核对测试记录并确认结果" in parsed.pages[1].raw_text

    table = next(block for block in parsed.blocks if block.block_type == "table")
    assert table.table_title == "表1 试验参数"
    assert table.table_header_rows == 1
    assert '<td colspan="2">设备应正常工作</td>' in (table.table_html or "")
    assert "记录1：参数=温度；下限=-20℃；上限=55℃" in table.text
    assert "记录2：参数=共同要求；下限=设备应正常工作；上限=设备应正常工作" in table.text

    body_chunk = next(chunk for chunk in parsed.chunks if "本文件规定" in chunk.text)
    assert body_chunk.section_path == ["无人机测试报告", "1 适用范围"]
    second_page_chunk = next(chunk for chunk in parsed.chunks if "逐项核对" in chunk.text)
    assert second_page_chunk.page_start == 2
    assert second_page_chunk.section_path == [
        "无人机测试报告",
        "1 适用范围",
        "1.1 验收方法",
    ]


def test_docx_parser_preserves_image_with_alt_text(tmp_path: Path) -> None:
    source = tmp_path / "带图片.docx"
    image_stream = BytesIO()
    Image.new("RGB", (40, 20), "navy").save(image_stream, format="PNG")
    image_stream.seek(0)
    document = Document()
    document.add_heading("结构说明", level=1)
    paragraph = document.add_paragraph()
    inline = paragraph.add_run().add_picture(image_stream)
    inline._inline.docPr.set("descr", "无人机系统结构示意图")
    document.save(source)

    parsed = DocxParser().parse(source)

    assert parsed.qa["image_count"] == 1
    assert parsed.qa["image_without_alt_count"] == 0
    assert len(parsed.assets) == 1
    assert parsed.assets[0].description == "无人机系统结构示意图"
    figure = next(block for block in parsed.blocks if block.block_type == "figure")
    assert figure.text == "无人机系统结构示意图"
    assert figure.asset_id == parsed.assets[0].asset_id


def test_docx_parser_rejects_renamed_zip(tmp_path: Path) -> None:
    source = tmp_path / "fake.docx"
    source.write_bytes(b"PK\x03\x04not-a-word-package")

    with pytest.raises(AppError) as error:
        DocxParser().parse(source)

    assert error.value.code == "DOCUMENT_SIGNATURE_MISMATCH"


def test_parser_registry_selects_docx_parser(tmp_path: Path) -> None:
    source = tmp_path / "sample.docx"

    parser = ParserRegistry.with_builtins().select(source)

    assert isinstance(parser, DocxParser)
