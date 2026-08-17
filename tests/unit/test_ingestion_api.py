from __future__ import annotations

import json
import time
from datetime import UTC, datetime
from io import BytesIO
from pathlib import Path

from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfWriter


def wait_for_terminal_job(client: TestClient, job_id: str) -> dict:
    deadline = time.monotonic() + 5
    while time.monotonic() < deadline:
        response = client.get(f"/api/v1/ingestion/jobs/{job_id}")
        job = response.json()
        if job["status"] not in {"queued", "running"}:
            return job
        time.sleep(0.02)
    raise AssertionError(f"Ingestion job did not finish in time: {job_id}")


def blank_pdf_bytes() -> bytes:
    stream = BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)
    writer.write(stream)
    return stream.getvalue()


def word_document_bytes() -> bytes:
    stream = BytesIO()
    document = Document()
    document.add_heading("测试要求", level=1)
    document.add_paragraph("无人机系统应完成上电自检并记录结果。")
    document.save(stream)
    return stream.getvalue()


def seed_editable_job(
    settings,
    *,
    job_id: str = "editable-job",
    source_name: str = "editable.pdf",
    source_relative_path: str | None = "法规/editable.pdf",
    publication_status: str = "planned",
    review_status: str = "unreviewed",
    created_at: str | None = None,
    with_table: bool = True,
) -> tuple[Path, Path]:
    root = Path(settings.ingestion.root_dir)
    source_sha256 = (job_id.encode("utf-8").hex() + "0" * 64)[:64]
    original = root / "originals" / source_sha256[:2] / source_sha256 / "source.pdf"
    original.parent.mkdir(parents=True, exist_ok=True)
    original.write_bytes(blank_pdf_bytes())
    output = root / "outputs" / job_id / "document"
    output.mkdir(parents=True, exist_ok=True)
    chunk = {
        "chunk_id": "chunk-1",
        "parent_chunk_id": "parent-1",
        "title": "第二十一条",
        "text": "旧的解析内容",
        "page_start": 2,
        "page_end": 2,
        "section_path": ["第二十一条"],
        "block_ids": ["block-1"],
        "keywords": ["旧内容"],
        "question_aliases": [],
        "asset_ids": [],
        "table_html": ["<table><tr><td>旧表格</td></tr></table>"] if with_table else [],
        "content_type": "table" if with_table else "text",
        "table_ids": ["table-1"] if with_table else [],
        "table_rows": [["旧表格"]] if with_table else [],
    }
    quality_gates = [{"gate": "text_coverage", "status": "pass", "message": "ok"}]
    document_ir = {
        "pages": [
            {
                "page_number": 2,
                "route": "native_text",
                "cleaned_text": "旧的解析内容",
                "cleaned_char_count": 7,
                "block_count": 1,
                "chunk_count": 1,
            }
        ],
        "blocks": [
            {
                "block_id": "block-1",
                "block_type": "paragraph",
                "text": "旧的解析内容",
                "page_number": 2,
                "source_page_start": 2,
                "source_page_end": 2,
            }
        ],
        "chunks": [chunk],
        "assets": [],
        "qa": {"quality_gates": quality_gates, "page_count": 1, "chunk_count": 1},
    }
    (output / "document_ir.json").write_text(
        json.dumps(document_ir, ensure_ascii=False),
        encoding="utf-8",
    )
    (output / "qa.json").write_text(
        json.dumps(document_ir["qa"], ensure_ascii=False),
        encoding="utf-8",
    )
    (output / "chunks.jsonl").write_text(
        json.dumps(chunk, ensure_ascii=False) + "\n",
        encoding="utf-8",
    )
    now = created_at or datetime.now(UTC).isoformat()
    job_path = root / "jobs" / job_id / "job.json"
    job_path.parent.mkdir(parents=True, exist_ok=True)
    job_path.write_text(
        json.dumps(
            {
                "job_id": job_id,
                "status": "completed",
                "source_name": source_name,
                "source_relative_path": source_relative_path,
                "source_sha256": source_sha256,
                "media_type": "application/pdf",
                "size_bytes": original.stat().st_size,
                "knowledge_base_id": "kb-regulations",
                "created_by": "tester",
                "original_path": str(original),
                "output_path": str(output),
                "quality": {
                    "route": "native_text",
                    "page_count": 1,
                    "chunk_count": 1,
                    "text_coverage": 1.0,
                    "warning_count": 0,
                    "failure_count": 0,
                },
                "publication": {
                    "status": publication_status,
                    "stage": publication_status,
                    "dataset_id": "kb-regulations",
                    "plan_sha256": "stale-plan",
                    "planned_chunk_count": 1,
                },
                "review": {
                    "status": review_status,
                    "reviewer": "reviewer-a" if review_status != "unreviewed" else None,
                    "updated_at": now,
                },
                "created_at": now,
                "updated_at": now,
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    return job_path, output


def test_ingestion_interface_and_assets_are_available(client: TestClient) -> None:
    interface = client.get("/ingestion")
    stylesheet = client.get("/assets/ingestion.css")
    script = client.get("/assets/ingestion.js")

    assert interface.status_code == 200
    assert 'id="fileInput"' in interface.text
    assert 'id="folderInput"' in interface.text
    assert "webkitdirectory" in interface.text
    assert 'id="selectFilesButton"' in interface.text
    assert 'id="selectFolderButton"' in interface.text
    assert 'id="jobList"' in interface.text
    assert 'id="jobSearchInput"' in interface.text
    assert 'id="clearJobSearch"' in interface.text
    assert 'class="queue-tabs"' in interface.text
    assert 'id="unreviewedJobCount"' in interface.text
    assert 'id="inReviewJobCount"' in interface.text
    assert 'id="approvedJobCount"' in interface.text
    assert 'id="reworkJobCount"' in interface.text
    assert 'id="publishedJobCount"' in interface.text
    assert 'id="selectAllApproved"' in interface.text
    assert 'id="batchPublishButton"' in interface.text
    assert 'id="batchPublishProgress"' in interface.text
    assert 'id="processingProgress"' in interface.text
    assert 'id="publicationProgress"' in interface.text
    assert 'id="reviewPanel"' in interface.text
    assert 'id="approveReviewNextButton"' in interface.text
    assert 'id="overrideReviewButton"' in interface.text
    assert 'id="reworkReviewNextButton"' in interface.text
    assert 'id="pageDiagnostics"' in interface.text
    assert 'id="pageDialog"' in interface.text
    assert 'id="pageAssets"' in interface.text
    assert 'id="chunkEditDialog"' in interface.text
    assert 'id="chunkEditText"' in interface.text
    assert 'id="firstIssuePageButton"' in interface.text
    assert stylesheet.status_code == 200
    assert "text/css" in stylesheet.headers["content-type"]
    assert script.status_code == 200
    assert "javascript" in script.headers["content-type"]
    assert "batchPublishSelected" in script.text
    assert "deleteJob" in script.text
    assert "openPageDetail" in script.text
    assert "openChunkEditor" in script.text
    assert "saveChunkEdit" in script.text
    assert "updateReviewStatus" in script.text
    assert "quality_gate_override" in script.text
    assert "loadAllJobs" in script.text
    assert 'method: "DELETE"' in script.text
    assert 'approved: "已通过资料"' in script.text
    assert 'published: "已发送资料"' in script.text


def test_job_search_matches_all_records_before_limit(
    client: TestClient,
    settings,
) -> None:
    seed_editable_job(
        settings,
        job_id="old-matching-job",
        source_name="2024无人机适航规章.pdf",
        source_relative_path="归档/适航/2024无人机适航规章.pdf",
    )
    seed_editable_job(
        settings,
        job_id="new-unrelated-job",
        source_name="unrelated.pdf",
        source_relative_path="其他/unrelated.pdf",
    )

    response = client.get(
        "/api/v1/ingestion/jobs",
        params={"limit": 1, "query": "适航"},
    )

    assert response.status_code == 200
    assert [job["job_id"] for job in response.json()] == ["old-matching-job"]


def test_review_workflow_persists_status_history_and_filters(
    client: TestClient,
    settings,
) -> None:
    job_path, _ = seed_editable_job(settings)

    started = client.patch(
        "/api/v1/ingestion/jobs/editable-job/review",
        json={"status": "in_review", "note": "正在核对跨页表格"},
        headers={"X-User-ID": "reviewer-a"},
    )
    assert started.status_code == 200
    assert started.json()["review"]["status"] == "in_review"
    assert started.json()["review"]["started_at"] is not None

    approved = client.patch(
        "/api/v1/ingestion/jobs/editable-job/review",
        json={"status": "approved", "note": "表格与条款均已核对"},
        headers={"X-User-ID": "reviewer-a"},
    )
    assert approved.status_code == 200
    assert approved.json()["review"]["status"] == "approved"
    assert approved.json()["review"]["reviewed_at"] is not None

    filtered = client.get(
        "/api/v1/ingestion/jobs",
        params={"review_status": "approved"},
    )
    assert filtered.status_code == 200
    assert [job["job_id"] for job in filtered.json()] == ["editable-job"]

    history = [
        json.loads(line)
        for line in (job_path.parent / "review_history.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
    ]
    assert [entry["to_status"] for entry in history] == ["in_review", "approved"]
    assert all(entry["reviewer"] == "reviewer-a" for entry in history)


def test_rework_review_requires_note(client: TestClient, settings) -> None:
    seed_editable_job(settings)

    response = client.patch(
        "/api/v1/ingestion/jobs/editable-job/review",
        json={"status": "rework"},
    )

    assert response.status_code == 422


def test_manual_confirmation_can_override_failed_quality_gate(
    client: TestClient,
    settings,
) -> None:
    job_path, output = seed_editable_job(settings)
    failed_gate = {
        "gate": "complex_table_fidelity",
        "status": "fail",
        "message": "需复核页面：[9, 15, 16]",
    }
    job = json.loads(job_path.read_text(encoding="utf-8"))
    job["status"] = "needs_review"
    job["quality"]["failure_count"] = 1
    job_path.write_text(json.dumps(job, ensure_ascii=False), encoding="utf-8")
    document_ir_path = output / "document_ir.json"
    document_ir = json.loads(document_ir_path.read_text(encoding="utf-8"))
    document_ir["qa"]["quality_gates"] = [failed_gate]
    document_ir_path.write_text(
        json.dumps(document_ir, ensure_ascii=False),
        encoding="utf-8",
    )
    (output / "qa.json").write_text(
        json.dumps(document_ir["qa"], ensure_ascii=False),
        encoding="utf-8",
    )

    blocked = client.patch(
        "/api/v1/ingestion/jobs/editable-job/review",
        json={"status": "approved", "note": "已经核对表格"},
        headers={"X-User-ID": "reviewer-a"},
    )
    assert blocked.status_code == 409
    assert blocked.json()["error"]["code"] == "INGESTION_REVIEW_QUALITY_BLOCKED"

    approved = client.patch(
        "/api/v1/ingestion/jobs/editable-job/review",
        json={
            "status": "approved",
            "note": "已人工修订并逐页核对",
            "quality_gate_override": True,
        },
        headers={"X-User-ID": "reviewer-a"},
    )
    assert approved.status_code == 200
    review = approved.json()["review"]
    assert review["status"] == "approved"
    assert review["quality_gate_override"] is True
    assert review["overridden_quality_gates"] == [failed_gate]
    assert review["quality_gate_overridden_at"] is not None

    history = [
        json.loads(line)
        for line in (job_path.parent / "review_history.jsonl")
        .read_text(encoding="utf-8")
        .splitlines()
    ]
    assert history[-1]["quality_gate_override"] is True
    assert history[-1]["overridden_quality_gates"] == [failed_gate]

    plan = client.post(
        "/api/v1/ingestion/jobs/editable-job/publish",
        json={"dataset_id": "kb-regulations", "dry_run": True},
    )
    assert plan.status_code == 200
    assert plan.json()["plan"]["chunk_count"] == 1

    edited = client.patch(
        "/api/v1/ingestion/jobs/editable-job/chunks/chunk-1",
        json={"text": "再次人工修订后的内容", "reason": "补充复核结果"},
        headers={"X-User-ID": "reviewer-a"},
    )
    assert edited.status_code == 200
    refreshed = client.get("/api/v1/ingestion/jobs/editable-job").json()
    assert refreshed["review"]["status"] == "unreviewed"
    assert refreshed["review"]["quality_gate_override"] is False
    assert refreshed["review"]["overridden_quality_gates"] == []
    assert refreshed["review"]["quality_gate_overridden_at"] is None


def test_quality_gate_override_only_applies_to_approval(
    client: TestClient,
    settings,
) -> None:
    seed_editable_job(settings)

    response = client.patch(
        "/api/v1/ingestion/jobs/editable-job/review",
        json={"status": "in_review", "quality_gate_override": True},
    )

    assert response.status_code == 422
    assert response.json()["error"]["code"] == "INGESTION_REVIEW_OVERRIDE_INVALID"


def test_job_pagination_uses_stable_upload_order(
    client: TestClient,
    settings,
) -> None:
    old_path, _ = seed_editable_job(
        settings,
        job_id="first-upload",
        source_name="z-last-name.pdf",
        created_at="2026-01-01T00:00:00+00:00",
    )
    seed_editable_job(
        settings,
        job_id="second-upload",
        source_name="a-first-name.pdf",
        created_at="2026-01-02T00:00:00+00:00",
    )
    old_payload = json.loads(old_path.read_text(encoding="utf-8"))
    old_payload["review"]["note"] = "修改文件不应改变首次上传顺序"
    old_path.write_text(json.dumps(old_payload, ensure_ascii=False), encoding="utf-8")

    first_page = client.get(
        "/api/v1/ingestion/jobs",
        params={"limit": 1, "offset": 0},
    )
    second_page = client.get(
        "/api/v1/ingestion/jobs",
        params={"limit": 1, "offset": 1},
    )

    assert first_page.status_code == 200
    assert second_page.status_code == 200
    assert first_page.json()[0]["job_id"] == "first-upload"
    assert second_page.json()[0]["job_id"] == "second-upload"


def test_manual_chunk_edit_persists_audit_and_invalidates_publish_plan(
    client: TestClient,
    settings,
) -> None:
    job_path, output = seed_editable_job(settings, review_status="approved")

    response = client.patch(
        "/api/v1/ingestion/jobs/editable-job/chunks/chunk-1",
        json={"text": "人工修正后的完整表格说明", "reason": "补全跨页表格"},
        headers={"X-User-ID": "reviewer-a"},
    )

    assert response.status_code == 200
    assert response.json()["revision"] == 1
    assert response.json()["content_type"] == "manual_text"

    document_ir = json.loads((output / "document_ir.json").read_text(encoding="utf-8"))
    chunk = document_ir["chunks"][0]
    assert chunk["text"] == "人工修正后的完整表格说明"
    assert chunk["table_html"] == []
    assert chunk["table_rows"] == []
    assert chunk["manual_revision"]["updated_by"] == "reviewer-a"
    assert chunk["manual_revision"]["reason"] == "补全跨页表格"
    assert chunk["manual_revision"]["replaced_structured_table"] is True
    assert document_ir["pages"][0]["manual_edited"] is True
    assert document_ir["qa"]["manual_revision_count"] == 1
    assert any(
        gate["gate"] == "manual_revision"
        for gate in document_ir["qa"]["quality_gates"]
    )
    chunks_jsonl = [
        json.loads(line)
        for line in (output / "chunks.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert chunks_jsonl[0]["text"] == "人工修正后的完整表格说明"
    audit = [
        json.loads(line)
        for line in (output / "manual_edits.jsonl").read_text(encoding="utf-8").splitlines()
    ]
    assert audit[0]["chunk_id"] == "chunk-1"
    assert audit[0]["revision"] == 1

    job = json.loads(job_path.read_text(encoding="utf-8"))
    assert job["status"] == "warning"
    assert job["quality"]["warning_count"] == 1
    assert job["publication"]["status"] == "not_requested"
    assert job["publication"]["plan_sha256"] is None
    assert job["publication"]["dataset_id"] == "kb-regulations"
    assert job["review"]["status"] == "unreviewed"
    assert "重新审核" in job["review"]["note"]

    detail = client.get("/api/v1/ingestion/jobs/editable-job/pages/2").json()
    assert detail["chunks"][0]["text"] == "人工修正后的完整表格说明"
    assert detail["chunks"][0]["manual_revision"]["revision"] == 1


def test_manual_chunk_edit_rejects_already_published_version(
    client: TestClient,
    settings,
) -> None:
    seed_editable_job(settings, publication_status="published")

    response = client.patch(
        "/api/v1/ingestion/jobs/editable-job/chunks/chunk-1",
        json={"text": "不应保存"},
    )

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "INGESTION_CHUNK_ALREADY_PUBLISHED"


def test_ingestion_queue_exposes_single_worker(client: TestClient) -> None:
    response = client.get("/api/v1/ingestion/queue")

    assert response.status_code == 200
    assert response.json() == {
        "worker_concurrency": 1,
        "active_job_id": None,
        "queued_job_ids": [],
    }


def test_upload_routes_empty_text_pdf_to_review(client: TestClient) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("scan.pdf", blank_pdf_bytes(), "application/pdf")},
        data={"knowledge_base_id": "regulations-test"},
        headers={"X-User-ID": "tester"},
    )

    assert response.status_code == 202
    job_id = response.json()["job_id"]

    job = wait_for_terminal_job(client, job_id)
    assert job["status"] == "needs_review"
    assert job["knowledge_base_id"] == "regulations-test"
    assert job["quality"]["route"] == "ocr_required"
    assert job["quality"]["chunk_count"] == 0
    assert job["progress"]["stage"] == "needs_review"
    assert job["progress"]["percent"] == 100

    preview_response = client.get(
        f"/api/v1/ingestion/jobs/{job_id}/preview",
        params={"chunk_limit": 5},
    )
    assert preview_response.status_code == 200
    preview = preview_response.json()
    assert preview["page_routes"] == {"blank_excluded": 1}
    assert preview["pages"][0]["page_number"] == 1
    assert preview["pages"][0]["route"] == "blank_excluded"
    assert preview["chunks"] == []

    page_response = client.get(f"/api/v1/ingestion/jobs/{job_id}/pages/1")
    assert page_response.status_code == 200
    page_detail = page_response.json()
    assert page_detail["page_number"] == 1
    assert page_detail["page"]["raw_text"] == ""
    assert page_detail["source_url"].endswith(f"/jobs/{job_id}/source#page=1")
    assert page_detail["assets"] == []

    output_folder = Path(job["output_path"])
    asset_folder = output_folder / "assets"
    asset_folder.mkdir()
    (asset_folder / "p0001-page-01.png").write_bytes(b"\x89PNG\r\n\x1a\n")
    document_ir_path = output_folder / "document_ir.json"
    document_ir = json.loads(document_ir_path.read_text(encoding="utf-8"))
    document_ir["assets"] = [
        {
            "asset_id": "p0001-page-01",
            "page_number": 1,
            "asset_type": "figure_page",
            "bbox": [0, 0, 595, 842],
            "filename": "page-0001.png",
            "mime_type": "image/png",
            "caption": "Page 1",
            "relative_path": "assets/p0001-page-01.png",
        }
    ]
    document_ir_path.write_text(
        json.dumps(document_ir, ensure_ascii=False),
        encoding="utf-8",
    )
    asset_detail = client.get(f"/api/v1/ingestion/jobs/{job_id}/pages/1").json()
    assert "?v=" in asset_detail["assets"][0]["url"]
    asset_response = client.get(asset_detail["assets"][0]["url"])
    assert asset_response.status_code == 200
    assert asset_response.headers["content-type"] == "image/png"
    assert asset_response.headers["content-disposition"] == "inline"
    assert asset_response.headers["cache-control"] == "private, no-store, max-age=0"

    source_response = client.get(f"/api/v1/ingestion/jobs/{job_id}/source")
    assert source_response.status_code == 200
    assert source_response.headers["content-type"].startswith("application/pdf")
    assert source_response.headers["content-disposition"] == "inline"


def test_upload_parses_docx_and_serves_canonical_media_type(client: TestClient) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={
            "file": (
                "测试要求.docx",
                word_document_bytes(),
                "application/octet-stream",
            )
        },
        headers={"X-User-ID": "tester"},
    )

    assert response.status_code == 202
    job = wait_for_terminal_job(client, response.json()["job_id"])
    assert job["status"] == "completed"
    assert job["media_type"] == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert job["quality"]["route"] == "docx_native"
    assert job["quality"]["chunk_count"] == 1

    source_response = client.get(
        f"/api/v1/ingestion/jobs/{job['job_id']}/source"
    )
    assert source_response.status_code == 200
    assert source_response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    preview = client.get(
        f"/api/v1/ingestion/jobs/{job['job_id']}/preview"
    ).json()
    assert preview["chunks"][0]["section_path"] == ["测试要求"]
    assert "上电自检" in preview["chunks"][0]["text"]


def test_folder_upload_preserves_safe_relative_path(client: TestClient) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("document.pdf", blank_pdf_bytes(), "application/pdf")},
        data={"source_relative_path": "法规/适航/../document.pdf"},
    )

    assert response.status_code == 202
    assert response.json()["source_relative_path"] is None

    accepted = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("document.pdf", blank_pdf_bytes(), "application/pdf")},
        data={"source_relative_path": r"法规\适航\document.pdf"},
    )

    assert accepted.status_code == 202
    assert accepted.json()["source_relative_path"] == "法规/适航/document.pdf"


def test_duplicate_upload_reuses_content_addressed_original(client: TestClient) -> None:
    document = blank_pdf_bytes()
    first = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("same.pdf", document, "application/pdf")},
    ).json()
    second = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("renamed.pdf", document, "application/pdf")},
    ).json()

    assert first["job_id"] != second["job_id"]
    assert first["source_sha256"] == second["source_sha256"]
    assert first["original_path"] == second["original_path"]


def test_retry_marks_previous_publication_as_stale(
    client: TestClient,
    settings,
) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("republish.pdf", blank_pdf_bytes(), "application/pdf")},
        data={"knowledge_base_id": "kb-existing"},
    )
    job_id = response.json()["job_id"]
    wait_for_terminal_job(client, job_id)
    job_path = (
        Path(settings.ingestion.root_dir)
        / "jobs"
        / job_id
        / "job.json"
    )
    persisted = json.loads(job_path.read_text(encoding="utf-8"))
    persisted["publication"] = {
        "status": "published",
        "stage": "published",
        "percent": 100,
        "message": "old result",
        "dataset_id": "kb-existing",
        "document_id": "old-document",
        "plan_sha256": "old-plan",
        "planned_chunk_count": 10,
        "published_chunk_count": 10,
    }
    job_path.write_text(
        json.dumps(persisted, ensure_ascii=False),
        encoding="utf-8",
    )

    retried = client.post(f"/api/v1/ingestion/jobs/{job_id}/retry")

    assert retried.status_code == 202
    publication = retried.json()["publication"]
    assert publication["status"] == "not_requested"
    assert publication["dataset_id"] == "kb-existing"
    assert publication["document_id"] is None
    assert publication["published_chunk_count"] == 0
    assert "尚未发送" in publication["message"]
    wait_for_terminal_job(client, job_id)


def test_rejects_unsupported_extension(client: TestClient) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("notes.txt", b"plain text", "text/plain")},
    )

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "UNSUPPORTED_DOCUMENT_TYPE"


def test_rejects_renamed_non_pdf_content(client: TestClient) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("fake.pdf", b"not a pdf", "application/pdf")},
    )

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "DOCUMENT_SIGNATURE_MISMATCH"


def test_rejects_renamed_non_docx_content(client: TestClient) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("fake.docx", b"not a word document", "application/octet-stream")},
    )

    assert response.status_code == 415
    assert response.json()["error"]["code"] == "DOCUMENT_SIGNATURE_MISMATCH"


def test_delete_unpublished_job_preserves_shared_original_until_last_reference(
    client: TestClient,
    settings,
) -> None:
    root = Path(settings.ingestion.root_dir)
    job_id = "delete-me"
    source_sha256 = "ab" * 32
    original = root / "originals" / source_sha256[:2] / source_sha256 / "source.pdf"
    output = root / "outputs" / job_id / "document"
    original.parent.mkdir(parents=True)
    original.write_bytes(blank_pdf_bytes())
    output.mkdir(parents=True)
    (output / "document_ir.json").write_text("{}", encoding="utf-8")
    now = datetime.now(UTC).isoformat()
    job_path = root / "jobs" / job_id / "job.json"
    job_path.parent.mkdir(parents=True)
    job_path.write_text(
        json.dumps(
            {
                "job_id": job_id,
                "status": "completed",
                "source_name": "delete.pdf",
                "source_sha256": source_sha256,
                "media_type": "application/pdf",
                "size_bytes": original.stat().st_size,
                "created_by": "tester",
                "original_path": str(original),
                "output_path": str(output),
                "created_at": now,
                "updated_at": now,
            }
        ),
        encoding="utf-8",
    )
    shared_job_id = "shared-copy"
    shared_job_path = root / "jobs" / shared_job_id / "job.json"
    shared_job_path.parent.mkdir(parents=True)
    shared_job_path.write_text(
        json.dumps(
            {
                "job_id": shared_job_id,
                "status": "failed",
                "source_name": "shared.pdf",
                "source_sha256": source_sha256,
                "media_type": "application/pdf",
                "size_bytes": original.stat().st_size,
                "created_by": "tester",
                "original_path": str(original),
                "created_at": now,
                "updated_at": now,
            }
        ),
        encoding="utf-8",
    )

    response = client.delete(f"/api/v1/ingestion/jobs/{job_id}")

    assert response.status_code == 204
    assert not job_path.parent.exists()
    assert not output.parent.exists()
    assert original.exists()
    assert client.get(f"/api/v1/ingestion/jobs/{job_id}").status_code == 404

    shared_response = client.delete(f"/api/v1/ingestion/jobs/{shared_job_id}")

    assert shared_response.status_code == 204
    assert not shared_job_path.parent.exists()
    assert not original.exists()


def test_delete_published_job_removes_only_local_ingestion_record(
    client: TestClient,
    settings,
) -> None:
    root = Path(settings.ingestion.root_dir)
    job_id = "published-job"
    source_sha256 = "cd" * 32
    original = root / "originals" / source_sha256[:2] / source_sha256 / "source.pdf"
    original.parent.mkdir(parents=True)
    original.write_bytes(blank_pdf_bytes())
    now = datetime.now(UTC).isoformat()
    job_path = root / "jobs" / job_id / "job.json"
    job_path.parent.mkdir(parents=True)
    job_path.write_text(
        json.dumps(
            {
                "job_id": job_id,
                "status": "completed",
                "source_name": "published.pdf",
                "source_sha256": source_sha256,
                "media_type": "application/pdf",
                "size_bytes": original.stat().st_size,
                "created_by": "tester",
                "original_path": str(original),
                "publication": {
                    "status": "published",
                    "dataset_id": "kb-1",
                    "document_id": "doc-1",
                },
                "created_at": now,
                "updated_at": now,
            }
        ),
        encoding="utf-8",
    )

    response = client.delete(f"/api/v1/ingestion/jobs/{job_id}")

    assert response.status_code == 204
    assert not job_path.exists()
    assert not original.exists()


def test_publish_dry_run_builds_plan_without_api_key(
    client: TestClient,
    settings,
) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("regulation.pdf", blank_pdf_bytes(), "application/pdf")},
        data={"knowledge_base_id": "regulations-test"},
    )
    job = response.json()
    wait_for_terminal_job(client, job["job_id"])
    job_path = (
        Path(settings.ingestion.root_dir)
        / "jobs"
        / job["job_id"]
        / "job.json"
    )
    persisted = json.loads(job_path.read_text(encoding="utf-8"))
    output = Path(settings.ingestion.root_dir) / "outputs" / job["job_id"] / "document"
    output.mkdir(parents=True)
    (output / "document_ir.json").write_text(
        json.dumps(
            {
                "chunks": [
                    {
                        "chunk_id": "chunk-1",
                        "title": "第一条",
                        "text": "测试内容",
                        "page_start": 1,
                        "page_end": 1,
                        "section_path": ["第一条"],
                    }
                ]
            },
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    persisted["status"] = "completed"
    persisted["output_path"] = str(output)
    job_path.write_text(json.dumps(persisted, ensure_ascii=False), encoding="utf-8")

    publish = client.post(
        f"/api/v1/ingestion/jobs/{job['job_id']}/publish",
        json={"dry_run": True},
    )

    assert publish.status_code == 200
    payload = publish.json()
    assert payload["dry_run"] is True
    assert payload["plan"]["chunk_count"] == 1
    assert payload["publication"]["status"] == "planned"
    assert payload["publication"]["stage"] == "planned"
    assert payload["publication"]["percent"] == 0


def test_real_publish_requires_manual_approval(
    client: TestClient,
    settings,
) -> None:
    seed_editable_job(settings, review_status="unreviewed")

    publish = client.post(
        "/api/v1/ingestion/jobs/editable-job/publish",
        json={"dataset_id": "kb-regulations", "dry_run": False},
    )

    assert publish.status_code == 409
    assert publish.json()["error"]["code"] == "INGESTION_REVIEW_REQUIRED"


def test_publish_rejects_quality_gate_failure(client: TestClient) -> None:
    response = client.post(
        "/api/v1/ingestion/jobs",
        files={"file": ("scan.pdf", blank_pdf_bytes(), "application/pdf")},
        data={"knowledge_base_id": "regulations-test"},
    )
    job_id = response.json()["job_id"]
    wait_for_terminal_job(client, job_id)

    publish = client.post(
        f"/api/v1/ingestion/jobs/{job_id}/publish",
        json={"dry_run": True},
    )

    assert publish.status_code == 409
    assert publish.json()["error"]["code"] == "INGESTION_QUALITY_GATE_BLOCKED"
